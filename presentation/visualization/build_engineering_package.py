#!/usr/bin/env python3
"""Build a genuine, verified Excel workbook from P001–P010 canonical CSV files.

No hand-authored XLSX XML, and no dependence on the previously broken workbook.
Usage: python presentation/visualization/build_engineering_package.py
"""
from __future__ import annotations

import csv
import io
import json
import math
import re
import unicodedata
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
VIS = ROOT / "presentation" / "visualization"
RES = VIS / "resources"
NAME = "JXwork_P001-P010_Engineering_QA.xlsx"
ZIPNAME = "JXwork_P001-P010_Engineering_QA_Pack.zip"
IDS = [f"P{i:03d}" for i in range(1, 11)]
EXPECTED = [13, 13, 9, 13, 12, 12, 16, 17, 18, 13]
WIDTHS = {"A": 22, "B": 43, "C": 78, "D": 17}
SOURCE_RE = re.compile(r"https?://[^\s，；\"<>\u3001]+")
DOCS = {
    "P002_AGENT_TOPICS.md": ("P002_AGENT_TOPICS.html", "专题说明"),
    "P005_NONNEGATIVE_RIDGE_GUIDE.md": ("P005_NONNEGATIVE_RIDGE_GUIDE.html", "算法图解"),
    "P006_BANDIT_COSTEER_GUIDE.md": ("P006_BANDIT_COSTEER_GUIDE.html", "算法图解"),
    "P010_SEMANTIC_PLAN_GUIDE.md": ("P010_SEMANTIC_PLAN_GUIDE.html", "流程图解"),
}

NAVY = "173650"
HEADER = "315575"
SLATE = "263E52"
WHITE = "FFFFFF"
ALT = "F2F6FA"
LINE = "DFE7EF"
BLUE = "146B9A"


def unicode_width(s: str) -> int:
    return sum(2 if unicodedata.east_asian_width(c) in ("F", "W") else 1 for c in s)


def line_count(s: str, width: int) -> int:
    # Wrapping is explicitly turned on. Use a shared, conservative height
    # estimate for EVERY sheet, not uncontrolled per-cell AutoFit.
    available = max(12, width - 6)
    lines = 0
    for paragraph in s.split("\n"):
        lines += max(1, math.ceil(unicode_width(paragraph) / available))
    return lines


def load_rows(pid: str) -> list[list[str]]:
    src = ROOT / "shared" / "papers" / pid / f"{pid}_工程复现问题清单_v1.csv"
    with src.open("r", encoding="utf-8-sig", newline="") as f:
        table = list(csv.reader(f))
    assert table and table[0] == ["主题", "关键问题", "简明答案"], (pid, "unexpected header")
    assert all(len(row) == 3 and all(v.strip() for v in row) for row in table[1:]), pid
    return table[1:]


def localize_link(pid: str, answer: str) -> tuple[str, str, str | None]:
    found = SOURCE_RE.findall(answer)
    if not found:
        return answer, "—", None
    if len(found) != 1:
        raise ValueError(f"{pid}: expected at most one external URL per answer: {found}")
    url = found[0].rstrip("。,.")
    if url.startswith("https://github.com/epshoron/JXwork/"):
        doc = next((name for name in DOCS if name in url), None)
        if doc:
            filename, label = DOCS[doc]
            target = f"resources/{filename}"
        else:
            raise ValueError(f"{pid}: unrecognized internal document URL {url}")
    else:
        target = f"resources/source_{pid}.html"
        label = "原始来源"
    if not (VIS / target).is_file():
        raise FileNotFoundError(f"Missing local hyperlink target: {target}")
    # Keep every word in the factual answer except the raw URL itself.
    cleaned = answer.replace(url, "").rstrip()
    cleaned = re.sub(r"(?:图解|详解|源码|原始仓库|原文|点击查看|入口)[：:]\s*$", "", cleaned)
    cleaned = cleaned.rstrip("：:；;, ")
    return cleaned, "↗ " + label, target


def make_workbook() -> tuple[bytes, list[dict], set[str]]:
    wb = Workbook()
    wb.remove(wb.active)
    report = []
    links_used: set[str] = set()
    thin = Side(style="hair", color=LINE)

    for pid, expected in zip(IDS, EXPECTED):
        data = load_rows(pid)
        if len(data) != expected:
            raise ValueError(f"{pid}: expected {expected} canonical questions, got {len(data)}")
        ws = wb.create_sheet(pid)
        ws.sheet_properties.tabColor = HEADER
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 95
        ws.freeze_panes = "A4"
        for letter, width in WIDTHS.items():
            ws.column_dimensions[letter].width = width

        ws.merge_cells("A1:D1")
        ws["A1"] = f"{pid}  /  工程复现问题"
        ws["A1"].font = Font(name="微软雅黑", size=15, color=WHITE, bold=True)
        ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 36
        # Fill entire title band explicitly, even when merged.
        for cell in ws[1][:4]:
            cell.fill = PatternFill("solid", fgColor=NAVY)

        ws.merge_cells("A2:D2")
        ws["A2"] = (
            f"{expected} 个核心问题    ·    十篇论文统一版式"
            "    ·    D 列点击查看本地补充资料"
        )
        ws["A2"].font = Font(name="微软雅黑", size=9, color="526B7B")
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[2].height = 26
        for col, header in enumerate(["主题", "关键问题", "简明答案", "补充资料"], 1):
            cell = ws.cell(3, col, header)
            cell.fill = PatternFill("solid", fgColor=HEADER)
            cell.font = Font(name="微软雅黑", size=10, color=WHITE, bold=True)
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[3].height = 30

        local_count = 0
        for i, (theme, question, raw_answer) in enumerate(data):
            rownum = i + 4
            answer, label, link = localize_link(pid, raw_answer)
            vals = [theme, question, answer, label]
            for col, value in enumerate(vals, 1):
                cell = ws.cell(rownum, col, value)
                cell.fill = PatternFill("solid", fgColor=WHITE if i % 2 == 0 else ALT)
                cell.border = Border(bottom=thin)
                cell.font = Font(name="微软雅黑", size=10, color=SLATE, bold=col == 1)
                cell.alignment = Alignment(
                    horizontal="left", vertical="center",
                    wrap_text=True, shrink_to_fit=False, indent=1
                )
            if link is not None:
                cell = ws.cell(rownum, 4)
                cell.hyperlink = link  # relative to XLSX, resolved inside ZIP package
                cell.font = Font(name="微软雅黑", size=10, color=BLUE, underline="single", bold=True)
                links_used.add(link)
                local_count += 1

            # SAME column widths / SAME row-height rule for all 10 sheets.
            # Leave enough lines to show answers, unlike previous hard-coded 70pt.
            lines = max(
                line_count(theme, WIDTHS["A"]),
                line_count(question, WIDTHS["B"]),
                line_count(answer, WIDTHS["C"]),
            )
            ws.row_dimensions[rownum].height = min(170, max(55, lines * 18 + 16))

        ws.auto_filter.ref = f"A3:D{3 + len(data)}"
        ws.sheet_format.defaultRowHeight = 18
        ws.print_options.horizontalCentered = True
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.print_title_rows = "1:3"
        report.append({"sheet": pid, "questions": expected, "links": local_count})

    wb.active = 0
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue(), report, links_used


def check_xlsx(data: bytes, expected_links: set[str]) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        assert zf.testzip() is None, "XLSX ZIP is corrupt"
    check = load_workbook(io.BytesIO(data), read_only=False, data_only=False)
    assert check.sheetnames == IDS, check.sheetnames
    total = 0
    seen_links = set()
    for pid, expected in zip(IDS, EXPECTED):
        ws = check[pid]
        assert ws.max_row == 3 + expected, (pid, ws.max_row)
        assert ws.freeze_panes == "A4"
        assert [ws.column_dimensions[c].width for c in WIDTHS] == list(WIDTHS.values())
        for r in range(4, 4 + expected):
            assert all(isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.strip() for c in (1, 2, 3))
            assert ws.cell(r, 3).alignment.wrap_text is True, (pid, r)
            assert ws.cell(r, 3).alignment.vertical == "center", (pid, r)
            assert ws.row_dimensions[r].height >= 55
            if ws.cell(r, 4).hyperlink is not None:
                seen_links.add(ws.cell(r, 4).hyperlink.target)
        total += expected
    assert total == 136
    assert seen_links == expected_links, (seen_links, expected_links)


def package(data: bytes, links: set[str], report: list[dict]) -> None:
    VIS.mkdir(parents=True, exist_ok=True)
    (VIS / NAME).write_bytes(data)
    zip_path = VIS / ZIPNAME
    resource_files = sorted(x for x in RES.iterdir() if x.is_file() and x.suffix in (".html", ".md"))
    assert len(resource_files) >= 14, "Expected four guide pairs + six source landing pages"
    assert (VIS / "README.html").exists()

    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=8) as z:
        z.writestr(NAME, data)
        z.write(VIS / "README.html", "README.html")
        for src in resource_files:
            z.write(src, "resources/" + src.name)
        z.writestr("README.txt",
            "P001–P010 工程问题整合包（重制版）\n"
            "请先完整解压，再打开 Excel；保持 resources/ 与 XLSX 同级。\n"
            "十张 Sheet 均用 Excel 文件库生成，答案自动换行、行高有足够空间。\n"
            "D 列链接打开本地 HTML；外部作者源站仍需联网。\n")
        z.writestr("manifest.json", json.dumps(
            {"sheets": IDS, "questions": 136, "linked_resources": sorted(links), "details": report},
            ensure_ascii=False, indent=2
        ))
    with zipfile.ZipFile(zip_path, "r") as z:
        assert z.testzip() is None
        assert NAME in z.namelist()
        assert all(target in z.namelist() for target in links)
        check_xlsx(z.read(NAME), links)
    print(json.dumps({
        "status": "OK", "sheets": 10, "questions": 136, "linked_files": len(links),
        "xlsx_bytes": (VIS / NAME).stat().st_size, "zip_bytes": zip_path.stat().st_size,
        "worksheets": report,
    }, ensure_ascii=False))


if __name__ == "__main__":
    workbook_bytes, sheet_report, hyperlink_targets = make_workbook()
    check_xlsx(workbook_bytes, hyperlink_targets)
    package(workbook_bytes, hyperlink_targets, sheet_report)
